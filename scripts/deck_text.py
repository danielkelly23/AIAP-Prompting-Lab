#!/usr/bin/env python3
"""Read a PowerPoint lecture as text: every slide's text as markdown, from the deck.

A week's lecture may be a PowerPoint deck, lectures-and-labs/weekNN/<topic>-lecture.pptx,
instead of a Marp deck (<topic>-lecture.md). The deck is the only source, and nothing
derived from it is committed: the site build (render_decks.py, build_index.py) and the
gates (verify_snippets.py, check_deck_portability.py) read it through this module when
they run. The decks carry no speaker notes, and any a deck still has are not read.

A code box says what it holds in its alt text, "Code, python" (the powerpoint-maker
skill writes it). Add ", no-parse" (right-click the box, View Alt Text) to exempt a
deliberately incomplete snippet from verify_snippets.py, as <!-- no-parse --> above
a fence does in a Marp deck.

Run it to read a deck, e.g. an assistant that cannot open a .pptx itself:
    python scripts/deck_text.py lectures-and-labs/week03/prompting-lecture.pptx
or with no argument to check every PowerPoint lecture: it renders each slide's
markdown as the site will and warns when a slide's words did not come through or some
markup shows as text (a gap in this script, never a reason to stop: the slides
themselves are fine). Runs anywhere, PowerPoint not needed
(pip install python-pptx markdown-it-py).
"""
from __future__ import annotations

import html
import io
import re
import sys
from itertools import groupby
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from schedule import load  # noqa: E402

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
MONO_FONTS = {"Consolas", "Cascadia Code", "Cascadia Mono", "Courier New", "Lucida Console"}


def escape(text: str) -> str:
    """Slide text as literal markdown: no emphasis, code or HTML it did not have."""
    text = text.replace("\\", "\\\\").replace("*", "\\*").replace("`", "\\`")
    text = re.sub(r"(?<![A-Za-z0-9])_|_(?![A-Za-z0-9])", r"\\_", text)
    return re.sub(r"<(?=[A-Za-z/!?])", "&lt;", text)


def escape_start(line: str) -> str:
    """Stop a paragraph that starts like a heading, list, quote or rule from becoming one."""
    if re.match(r"#{1,6}\s|>|[-+*]\s|(-{3,}|\*{3,}|_{3,}|=+)\s*$", line):
        return "\\" + line
    m = re.match(r"\d+(?=[.)]\s)", line)
    return line[:m.end()] + "\\" + line[m.end():] if m else line


def code_span(text: str) -> str:
    ticks = "`"
    while ticks in text:
        ticks += "`"
    pad = " " if text.startswith("`") or text.endswith("`") else ""
    return f"{ticks}{pad}{text}{pad}{ticks}"


def wrap(marks: str, text: str) -> str:
    """`marks` around text, whitespace kept outside: `** x**` is not bold in markdown."""
    core = text.strip()
    if not marks or not core:
        return text
    return f"{text[:len(text) - len(text.lstrip())]}{marks}{core}{marks}{text[len(text.rstrip()):]}"


def line_markdown(runs: list[tuple]) -> str:
    """One line of (text, bold, italic, mono, link) runs as markdown. Emphasis
    NESTS rather than being set run by run: bold, then bold italic, then bold
    again must come out as **a *b* c**, because **a** ***b***** c** leaves a
    run of five asterisks that markdown cannot pair and prints."""
    out = []
    for link, same_link in groupby(runs, key=lambda r: r[4]):
        text = ""
        for bold, same_bold in groupby(same_link, key=lambda r: r[1]):
            inner = ""
            for italic, same in groupby(same_bold, key=lambda r: r[2]):
                inner += wrap("*" if italic else "",
                              "".join(code_run(t) if mono else escape(t)
                                      for t, _, _, mono, _ in same))
            text += wrap("**" if bold else "", inner)
        out.append(f"[{text.strip()}]({link})" if link else text)
    return "".join(out)


def code_run(text: str) -> str:
    """A monospace run as a code span, its outer spaces left outside the backticks."""
    if not text.strip():
        return text
    return (text[:len(text) - len(text.lstrip())] + code_span(text.strip())
            + text[len(text.rstrip()):])


def inline(paragraph, mono=True) -> str:
    """A paragraph's runs as inline markdown: **bold**, *italic*, `code`, links.
    A line break inside the paragraph comes back as a newline. mono=False reads
    Consolas as plain text: a prompt, a table header or a "//" remark is set in
    Consolas without being code."""
    from pptx.text.text import _Run
    lines: list[list[tuple]] = [[]]
    for el in paragraph._p:
        if el.tag == A + "r":
            run = _Run(el, paragraph)
            link = run.hyperlink.address if el.find(f"{A}rPr/{A}hlinkClick") is not None else None
            lines[-1].append((run.text, bool(run.font.bold), bool(run.font.italic),
                              mono and run.font.name in MONO_FONTS, link))
        elif el.tag == A + "br":
            lines.append([])
        elif el.tag == A + "fld":
            lines[-1].append(("".join(t.text or "" for t in el.iter(A + "t")),
                              False, False, False, None))
    # A card's label ("THEN · 2023", "HOW YOU CAN TELL") is Consolas capitals
    # from end to end: it names the card, it is not code.
    printed = [r for line in lines for r in line if r[0].strip()]
    text = "".join(r[0] for r in printed)
    if printed and all(r[3] for r in printed) and re.search(r"[A-Z]", text) and text == text.upper():
        lines = [[(t, b, i, False, link) for t, b, i, _, link in line] for line in lines]
    return "\n".join(line_markdown(runs).strip() for runs in lines).strip()


def cnvpr(shape):
    return shape._element.find(f".//{P}cNvPr")


def layout_list_style(shape, slide):
    """A placeholder's list style as its layout defines it, or None."""
    if not shape.is_placeholder:
        return None
    base = slide.slide_layout.placeholders.get(idx=shape.placeholder_format.idx)
    return base._element.find(f".//{A}lstStyle") if base is not None else None


def list_styles(shape, slide) -> list:
    """Where a paragraph's bullet setting is inherited from: the shape's own
    list style, then (for a placeholder) its layout placeholder's."""
    styles = [shape._element.find(f".//{A}lstStyle"), layout_list_style(shape, slide)]
    return [s for s in styles if s is not None]


def bullet_setting(ppr) -> bool | None:
    if ppr is None:
        return None
    if ppr.find(A + "buNone") is not None:
        return False
    if ppr.find(A + "buChar") is not None or ppr.find(A + "buAutoNum") is not None:
        return True
    return None


def is_bullet(paragraph, shape, slide) -> bool:
    setting = bullet_setting(paragraph._p.find(A + "pPr"))
    for style in list_styles(shape, slide):
        if setting is None:
            setting = bullet_setting(style.find(f"{A}lvl{paragraph.level + 1}pPr"))
    if setting is not None:
        return setting
    # Otherwise the master decides: body and content placeholders are bulleted.
    return shape.is_placeholder and shape.placeholder_format.type in (2, 7)   # BODY, OBJECT


def is_heading(shape, slide) -> bool:
    """A placeholder its layout sets bold and unbulleted, e.g. a Comparison
    slide's column headings."""
    style = layout_list_style(shape, slide)
    lvl1 = style.find(A + "lvl1pPr") if style is not None else None
    if lvl1 is None or lvl1.find(A + "buNone") is None:
        return False
    rpr = lvl1.find(A + "defRPr")
    return rpr is not None and rpr.get("b") in ("1", "true")


def text_blocks(shape, slide) -> list[str]:
    heading = is_heading(shape, slide)
    blocks, items = [], []
    for paragraph in shape.text_frame.paragraphs:
        plain = "".join(r.text for r in paragraph.runs).strip()
        md = escape(plain) if plain.startswith("//") else inline(paragraph)   # a "// remark"
        if not md:
            continue
        if heading:
            blocks.append("### " + md.replace("\n", " "))
        elif is_bullet(paragraph, shape, slide):
            indent = "  " * paragraph.level
            items.append(f"{indent}- " + escape_start(md).replace("\n", f"\\\n{indent}  "))
            continue
        else:
            blocks.append(escape_start(md).replace("\n", "\\\n"))
        if items:                              # a list ends where anything else starts
            blocks.insert(len(blocks) - 1, "\n".join(items))
            items = []
    if items:
        blocks.append("\n".join(items))
    return blocks


def code_block(shape) -> str:
    """A code box as a fence, its language (and any no-parse) from its alt text."""
    words = [w.strip().lower() for w in (cnvpr(shape).get("descr") or "").split(",")]
    lang = words[1] if len(words) > 1 and words[1] else "text"
    lang = "text" if lang == "plain text" else re.sub(r"[^a-z0-9_+-]", "", lang)
    lines = []
    for paragraph in shape.text_frame.paragraphs:
        line = "".join("\n" if el.tag == A + "br" else "".join(t.text or "" for t in el.iter(A + "t"))
                       for el in paragraph._p if el.tag in (A + "r", A + "br", A + "fld"))
        lines.append(line.rstrip())
    while lines and not lines[-1]:
        lines.pop()
    code = "\n".join(lines)
    fence = "```"
    while fence in code:
        fence += "`"
    block = f"{fence}{lang}\n{code}\n{fence}"
    return ("<!-- no-parse -->\n" + block) if "no-parse" in words[2:] else block


def diagram_block(group, slide) -> str:
    """A flow as a numbered list, left to right; a stack as a list, top to bottom.
    A stack's items take `+`, not `-`: in CommonMark two `-` lists with only a
    blank line between them are ONE list, and the slide's bullets often follow."""
    boxes = [s for s in group.shapes if s.has_text_frame and s.text_frame.text.strip()]
    if group.name.startswith("Flow"):
        items = []
        for n, box in enumerate(sorted(boxes, key=lambda s: s.left), 1):
            paras = [inline(p) for p in box.text_frame.paragraphs if inline(p)]
            if len(paras) > 1 and re.fullmatch(r"\**\d{1,2}\**", paras[0]):
                paras = paras[1:]              # the step number the box prints
            items.append(f"{n}. " + escape_start(" ".join(paras)).replace("\n", " "))
        return "\n".join(items)
    items = []
    for box in sorted(boxes, key=lambda s: s.top):
        md = " ".join(inline(p) for p in box.text_frame.paragraphs if inline(p)).replace("\n", " ")
        text, _, label = md.partition("\t")
        label = label.strip().strip("*").strip()
        items.append("+ " + escape_start(text.strip()) + (f" · *{label}*" if label else ""))
    return "\n".join(items)


def table_block(shape) -> str:
    def cell(c, header=False):
        return " ".join(inline(p, mono=not header) for p in c.text_frame.paragraphs).replace("|", "\\|").replace("\n", " ")
    rows = [[cell(c, r == 0) for c in row.cells] for r, row in enumerate(shape.table.rows)]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
    return "\n".join(lines + ["| " + " | ".join(r) + " |" for r in rows[1:]])


BOXES = ("Prompt ", "Reply ", "Callout ")


def box_block(shape) -> str:
    """A prompt, reply or callout box as a blockquote, its label in bold. A prompt's
    Consolas is how it is typed, not code, so it comes back as plain text."""
    prompt = shape.name.startswith("Prompt ")
    labelled = shape.name.startswith(("Prompt ", "Reply "))       # PROMPT / REPLY, in Consolas
    lines = [inline(p, mono=not (prompt or (labelled and k == 0)))
             for k, p in enumerate(shape.text_frame.paragraphs)]
    lines = [escape_start(line).replace("\n", chr(92) + "\n> ") for line in lines if line]
    return "\n>\n".join("> " + line for line in lines)


def shape_blocks(shape, slide, warnings: list[str], where: str) -> list[str]:
    descr = (cnvpr(shape).get("descr") or "").strip() if cnvpr(shape) is not None else ""
    if shape.shape_type == 6:                  # a group
        if shape.name.startswith(("Flow", "Stack")):
            return [diagram_block(shape, slide)]
        return [b for s in shape.shapes for b in shape_blocks(s, slide, warnings, where)]
    if descr.lower().startswith("code") and shape.has_text_frame:
        return [code_block(shape)]
    if getattr(shape, "has_table", False) and shape.has_table:
        return [table_block(shape)]
    if shape.has_text_frame and shape.name.startswith(BOXES):
        return [box_block(shape)]
    if shape.has_text_frame:
        return text_blocks(shape, slide)
    if shape.shape_type == 13:                 # a picture
        return [f"*\\[Picture: {escape(descr)}\\]*" if descr else "*\\[Picture\\]*"]
    if getattr(shape, "has_chart", False) and shape.has_chart and descr.startswith("Chart:"):
        return [f"*\\[{escape(descr)}\\]*"]         # the builder writes the data into the alt text
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return [f"*\\[Chart: {escape(descr)}\\]*" if descr else "*\\[Chart\\]*"]
    if shape._element.tag == P + "graphicFrame":
        warnings.append(f"{where}: a SmartArt or other object's text cannot be read as text")
        return [f"*\\[Diagram: {escape(descr)}\\]*" if descr else "*\\[Diagram\\]*"]
    return []


def slides_of(deck_bytes: bytes, warnings: list[str]) -> tuple[str, list[str]]:
    """The deck's title, and each slide shown (hidden ones are left out) as markdown."""
    from pptx import Presentation
    prs = Presentation(io.BytesIO(deck_bytes))
    title, out = "", []
    for slide in prs.slides:
        if slide._element.get("show") == "0":
            continue
        where = f"slide {len(out) + 1}"
        head = slide.shapes.title
        blocks = []
        if head is not None and head.text_frame.text.strip():
            text = " ".join(inline(p) for p in head.text_frame.paragraphs if inline(p)).replace("\n", " ")
            blocks.append(("#" if not out else "##") + " " + text)
            title = title or head.text_frame.text.replace("\v", " ").replace("\n", " ").strip()
        for shape in slide.shapes:
            if head is None or shape.shape_id != head.shape_id:
                blocks += shape_blocks(shape, slide, warnings, where)
        out.append("\n\n".join(blocks))
    return title or prs.core_properties.title, out


def deck_texts(deck_bytes: bytes) -> list[str]:
    """Each shown slide's text as the deck holds it, for round_trip()."""
    from pptx import Presentation
    out = []
    for slide in Presentation(io.BytesIO(deck_bytes)).slides:
        if slide._element.get("show") == "0":
            continue
        texts = []

        def walk(shapes):
            for sh in shapes:
                if sh.shape_type == 6:
                    walk(sh.shapes)
                elif sh.has_text_frame:
                    for p in sh.text_frame.paragraphs:
                        text = "".join(r.text for r in p.runs)
                        if not (sh.name.startswith("Step") and re.fullmatch(r"\s*\d{1,2}\s*", text)):
                            texts.append(text)
                elif getattr(sh, "has_table", False) and sh.has_table:
                    texts.extend(c.text for row in sh.table.rows for c in row.cells)
                else:
                    descr = (cnvpr(sh).get("descr") or "").strip() if cnvpr(sh) is not None else ""
                    if sh.shape_type == 13:
                        texts.append(f"Picture: {descr}" if descr else "Picture")
                    elif getattr(sh, "has_chart", False) and sh.has_chart:
                        texts.append(descr if descr.startswith("Chart:") else f"Chart: {descr}")
        walk(slide.shapes)
        out.append(" ".join(texts))
    return out


def round_trip(deck_bytes: bytes, slides: list[str]) -> list[str]:
    """Render each slide's markdown as the site will (markdown-it) and compare
    its words with the deck's: nothing lost, nothing added, and no `*` or
    backtick left over, which is markup that failed to parse."""
    from build_lab_pages import RENDERER
    problems = []
    for n, (source, md) in enumerate(zip(deck_texts(deck_bytes), slides), 1):
        shown = html.unescape(re.sub(r"<[^>]+>", " ", RENDERER.render(
            re.sub(r"<!--.*?-->", "", md, flags=re.S))))
        if re.findall(r"\w+", shown) != re.findall(r"\w+", source):
            problems.append(f"slide {n}: its words do not survive the conversion")
        for mark in "*`":
            if shown.count(mark) != source.count(mark):
                problems.append(f"slide {n}: a {mark!r} shows as text, so some markup did not parse")
    return problems


def read_deck(path: Path) -> bytes:
    try:
        return path.read_bytes()
    except PermissionError:
        raise SystemExit(f"{path.as_posix()} is locked: close it in PowerPoint and run this again")


def deck_title(path: Path) -> str:
    """The deck's title: its first slide's title, else its document title."""
    return slides_of(read_deck(path), [])[0]


def main() -> None:
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(encoding="utf-8", errors="replace")
    if len(sys.argv) > 1:                       # print each named deck's text
        for name in sys.argv[1:]:
            title, slides = slides_of(read_deck(Path(name)), [])
            print("\n\n---\n\n".join(slides))
        return
    decks = [r for r in load().teaching if r.is_pptx]
    for row in decks:                           # check every PowerPoint lecture
        deck_bytes = read_deck(row.lecture)
        warnings: list[str] = []
        title, slides = slides_of(deck_bytes, warnings)
        warnings += round_trip(deck_bytes, slides)
        codes = sum(len(re.findall(r"(?m)^`{3,}[a-z]", md)) for md in slides)
        print(f"{row.lecture.as_posix()}: {title!r}, {len(slides)} slides, {codes} code boxes")
        for w in warnings:
            print(f"  warning: {w}")
    print(f"deck_text: {len(decks)} PowerPoint deck(s) read")


if __name__ == "__main__":
    main()
